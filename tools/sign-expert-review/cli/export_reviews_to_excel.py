#!/usr/bin/env python3
"""Copy a MASTER workbook and append a timestamped portal-result sheet."""

from __future__ import annotations

import argparse
import json
import shutil
from datetime import datetime, timezone
from pathlib import Path

from review_common import ReviewBlockedError, read_json, reduce_latest_events, validate_manifest


HEADERS = (
    "후보 번호", "signer ID", "검토 revision", "검토자 코드", "검토 시각",
    "검토 모드", "검토 상태", "글로스 검증", "경계 방향성 평가", "정밀 검토 필요",
    "정밀 검토 사유", "기존 시작 frame", "기존 종료 frame", "수동 시작 frame",
    "수동 종료 frame", "승인 시작 frame", "승인 종료 frame", "승인 시작 source sec",
    "승인 종료 source sec", "앱 학습 영상", "내부 reference", "대표 영상 적합도",
    "문장 연결 영향", "제외 여부", "제외 사유", "전문가 의견", "검토 시간 ms", "클릭 수",
)


def export_excel(workbook: Path, manifest: dict, current_document: dict, output_dir: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ReviewBlockedError("EXCEL_EXPORT_DEPENDENCY_MISSING", "openpyxl is required for Excel export") from exc
    if not workbook.exists() or workbook.suffix.casefold() != ".xlsx":
        raise ReviewBlockedError("EXCEL_EXPORT_INVALID", "input workbook must be an existing .xlsx file")
    if current_document.get("batch_id") != manifest["batch_id"] or current_document.get("revision") != manifest["revision"]:
        raise ReviewBlockedError("STALE_MANIFEST", "review_current does not match the batch revision")
    events, issues = reduce_latest_events(current_document.get("events") or [], manifest)
    if issues:
        raise ReviewBlockedError("INVALID_USAGE_COMBINATION", "review_current contains invalid events", details={"issues": issues})
    candidates = {candidate["candidate_id"]: candidate for candidate in manifest["candidates"]}
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / f"{workbook.stem}_portal_export_{timestamp}.xlsx"
    shutil.copy2(workbook, output)
    book = load_workbook(output)
    title = f"portal_{timestamp[:8]}"
    index = 1
    while title in book.sheetnames:
        index += 1
        title = f"portal_{timestamp[:8]}_{index}"
    sheet = book.create_sheet(title)
    sheet.append(HEADERS)
    for event in events:
        candidate = candidates[event["candidate_id"]]
        sheet.append((
            event["candidate_id"], candidate["signer_id"], event["revision"], event["reviewer_code"], event["reviewed_at"],
            event["review_mode"], event["review_state"], event["target_validation"], event["boundary_assessment"],
            event["precise_review_required"], event["precise_review_reason"], candidate["annotated_start_local_frame"],
            candidate["annotated_end_local_frame"], event["manual_start_local_frame"], event["manual_end_local_frame"],
            event["approved_start_local_frame"], event["approved_end_local_frame"], event["approved_start_source_sec"],
            event["approved_end_source_sec"], event["learning_video_assessment"], event["reference_assessment"],
            event["representative_fit"], event["sentence_connection_effect"], event["excluded"],
            event["exclusion_reason"], event["expert_notes"], event["review_duration_ms"], event["interaction_count"],
        ))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    book.save(output)
    report = output.with_suffix(".changes.md")
    report.write_text(
        "\n".join([
            "# Portal Excel export change report", "", f"- Source workbook: `{workbook.name}`",
            f"- Output workbook: `{output.name}`", f"- Added sheet: `{title}`",
            f"- Exported candidates: `{len(events)}`", "- Source workbook overwritten: `false`", "",
        ]),
        encoding="utf-8",
    )
    return {"status": "EXCEL_EXPORT_COMPLETE", "output": str(output), "change_report": str(report), "exported_count": len(events)}


def parser() -> argparse.ArgumentParser:
    value = argparse.ArgumentParser(description=__doc__)
    value.add_argument("--workbook", type=Path, required=True)
    value.add_argument("--review-manifest", type=Path, required=True)
    value.add_argument("--review-current", type=Path, required=True)
    value.add_argument("--output-dir", type=Path, required=True)
    return value


def main() -> int:
    args = parser().parse_args()
    try:
        report = export_excel(
            args.workbook, validate_manifest(read_json(args.review_manifest)),
            read_json(args.review_current), args.output_dir,
        )
    except ReviewBlockedError as exc:
        print(json.dumps(exc.as_dict(), ensure_ascii=False, indent=2))
        return 2
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
